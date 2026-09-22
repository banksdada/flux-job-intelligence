"""
F1: CV Upload & Parsing.

Extracts plain text from an uploaded .docx / .txt / .md CV, identifies
skills/certifications and target roles from it, and estimates years of
experience — everything downstream (scoring, filtering) works off the
resulting evidence set. Nothing here leaves the machine it runs on.
"""
import io
import re

import docx

from skills_data import detect_roles, estimate_years_experience, find_skills

_MD_MARKUP = re.compile(r"[#*_`>\[\]\(\)-]")


def extract_text(filename: str, file_bytes: bytes) -> str:
    """Dispatch on extension and return the CV's plain text content."""
    lower = filename.lower()
    if lower.endswith(".docx"):
        return _extract_docx(file_bytes)
    if lower.endswith(".md"):
        return _extract_markdown(file_bytes)
    if lower.endswith(".txt"):
        return file_bytes.decode("utf-8", errors="ignore")
    raise ValueError(f"Unsupported CV file type: {filename}")


def _extract_docx(file_bytes: bytes) -> str:
    document = docx.Document(io.BytesIO(file_bytes))
    parts = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    parts.append(cell.text)
    return "\n".join(parts)


def _extract_markdown(file_bytes: bytes) -> str:
    raw = file_bytes.decode("utf-8", errors="ignore")
    return _MD_MARKUP.sub(" ", raw)


def parse_cv(filename: str, file_bytes: bytes) -> dict:
    """
    Returns:
        {
          "text": full extracted text (kept for session re-scoring),
          "evidence": sorted list of canonical skill/cert labels found,
          "roles": ordered list of detected role keys (strongest first),
          "filter_role": best-guess role key to auto-select, or None,
          "years_experience": int estimate,
        }
    """
    text = extract_text(filename, file_bytes)
    evidence = sorted(find_skills(text))
    roles = detect_roles(text)
    years = estimate_years_experience(text)
    return {
        "text": text,
        "evidence": evidence,
        "roles": roles,
        "filter_role": roles[0] if roles else None,
        "years_experience": years,
    }
